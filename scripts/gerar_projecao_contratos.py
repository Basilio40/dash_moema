"""
Gera src/data/projecao-contratos.json a partir do relatório Focco de contratos
(loja_18314_MOEMA_CONTRATOS/relatorios_focco), agregando por mês de
"Previsão de Entrega":
  - valor: soma de "Valor da Venda" do mês (em REAIS)
  - projetos: nº de contratos com entrega prevista no mês (base da medição)

O eixo X do gráfico "Previsão de Entradas × Saídas — com Projeção de Despesas
com Vendas" passa a ser exatamente a sequência de meses desta série.
"""
import glob
import json
import os
import re
import sys
import warnings
from datetime import datetime

import openpyxl

warnings.filterwarnings("ignore")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTRATOS_DIR = os.path.join(
    os.path.dirname(BASE_DIR), 'loja_18314_MOEMA_CONTRATOS', 'relatorios_focco')

MESES_PT = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun',
            'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']


def xlsx_mais_recente(diretorio):
    arquivos = [f for f in glob.glob(os.path.join(diretorio, '*.xlsx'))
                if not os.path.basename(f).startswith('~$')]
    if not arquivos:
        raise FileNotFoundError(f'Nenhum .xlsx encontrado em {diretorio}')
    return max(arquivos, key=os.path.getmtime)


def achar_cabecalho(ws):
    """Linha do cabeçalho da tabela de contratos (colunas Contrato, Valor da
    Venda, Previsão de Entrega...). O relatório Focco traz filtros antes."""
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        celulas = [str(c).strip() if c is not None else '' for c in row]
        if 'Valor da Venda' in celulas and 'Previsão de Entrega' in celulas:
            return i, {nome: j for j, nome in enumerate(celulas) if nome}
    raise ValueError('Cabeçalho "Valor da Venda"/"Previsão de Entrega" não encontrado')


def parse_valor(celula):
    """'R$ 20.146,00' / '20146.00' / 20146.00 -> float (reais)."""
    if celula is None:
        return 0.0
    if isinstance(celula, (int, float)):
        return float(celula)
    txt = re.sub(r'[^\d,.-]', '', str(celula))
    if not txt:
        return 0.0
    if ',' in txt:  # pt-BR: 20.146,00
        txt = txt.replace('.', '').replace(',', '.')
    return float(txt)


def parse_data(celula):
    if isinstance(celula, datetime):
        return celula
    if celula is None:
        return None
    try:
        return datetime.strptime(str(celula).strip(), '%d/%m/%Y')
    except ValueError:
        return None


def rotulo_mes(d):
    return f'{MESES_PT[d.month - 1]}/{str(d.year)[2:]}'


def main():
    caminho = xlsx_mais_recente(CONTRATOS_DIR)
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb.worksheets[0]

    linha_header, colunas = achar_cabecalho(ws)
    idx_valor = colunas['Valor da Venda']
    idx_entrega = colunas['Previsão de Entrega']

    por_mes = {}  # (ano, mes) -> {'valor': x, 'projetos': n}
    total_geral = 0.0
    n_contratos = 0
    n_sem_data = 0

    for row in ws.iter_rows(min_row=linha_header + 1, values_only=True):
        entrega = parse_data(row[idx_entrega] if idx_entrega < len(row) else None)
        valor = parse_valor(row[idx_valor] if idx_valor < len(row) else None)
        if entrega is None:
            if valor:
                n_sem_data += 1
            continue
        chave = (entrega.year, entrega.month)
        agg = por_mes.setdefault(chave, {'valor': 0.0, 'projetos': 0})
        agg['valor'] += valor
        agg['projetos'] += 1
        total_geral += valor
        n_contratos += 1

    acumulado = 0.0
    series = []
    for (ano, mes) in sorted(por_mes):
        agg = por_mes[(ano, mes)]
        acumulado += agg['valor']
        series.append({
            'mes': rotulo_mes(datetime(ano, mes, 1)),
            'valor': round(agg['valor'], 2),
            'projetos': agg['projetos'],
            'acumulado': round(acumulado, 2),
        })

    saida = {
        'fonte': os.path.basename(caminho),
        'colunas': ['Valor da Venda', 'Previsão de Entrega'],
        'total': round(total_geral, 2),
        'numContratos': n_contratos,
        'semPrevisaoEntrega': n_sem_data,
        'series': series,
    }

    destino = os.path.join(BASE_DIR, 'src', 'data', 'projecao-contratos.json')
    with open(destino, 'w', encoding='utf-8') as f:
        json.dump(saida, f, ensure_ascii=False, indent=2)

    print(f'projecao-contratos: {len(series)} meses, '
          f'{n_contratos} contratos, total R$ {total_geral:,.2f}'
          .replace(',', 'X').replace('.', ',').replace('X', '.'))
    if n_sem_data:
        print(f'  aviso: {n_sem_data} contratos com valor mas sem previsão de entrega')


if __name__ == '__main__':
    sys.exit(main())
