# -*- coding: utf-8 -*-
"""Gera src/data/fluxo-caixa-realizado.json a partir dos relatorios de realizado
da loja 18314 (loja_18314_FLUXO_DE_CAIXA/relatorios_fluxo_caixa/realizado).

Soma as colunas Credito (entradas) e Debito (saidas) de cada
Realizado_<Mes>_<Ano>.xlsx e mantem apenas os meses do ano atual com movimento.

Atencao: o nome do arquivo nem sempre bate com o periodo real do relatorio
(ex.: Realizado_Fev_2026.xlsx veio com periodo 01/01/2026 ate 31/01/2026),
por isso o mes/ano e derivado da linha "Periodo de Lancamento" dentro do
proprio xlsx; o nome do arquivo e usado apenas como fallback.

Uso: python scripts/gerar_realizado_fluxo.py
"""
import glob
import json
import os
import re
import warnings
from datetime import date

import openpyxl

warnings.filterwarnings("ignore")

PASTA_REALIZADO = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "..",
    "loja_18314_FLUXO_DE_CAIXA",
    "relatorios_fluxo_caixa",
    "realizado",
)
SAIDA = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "src",
    "data",
    "fluxo-caixa-realizado.json",
)

ORDEM = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun",
         "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
INDICE_MES = {m: i + 1 for i, m in enumerate(ORDEM)}


def colunas_credito_debito(ws):
    """Localiza a linha de cabecalho e devolve os indices das colunas
    Credito e Debito (o layout do realizado tem colunas a mais que o
    de previsao, entao nao da para fixar posicoes)."""
    for linha in ws.iter_rows(values_only=True):
        celulas = [str(c).strip().lower() if c is not None else "" for c in linha]
        if "crédito" in celulas and "débito" in celulas:
            return celulas.index("crédito"), celulas.index("débito")
        if "credito" in celulas and "debito" in celulas:
            return celulas.index("credito"), celulas.index("debito")
    return 7, 8  # layout padrao do relatorio realizado


def mes_ano_do_periodo(ws):
    """Extrai (mes, ano) da data inicial da linha 'Periodo de ...'.
    Retorna None se o campo estiver vazio ou ilegivel."""
    for linha in ws.iter_rows(min_row=1, max_row=12, values_only=True):
        for celula in linha[:3]:
            texto = str(celula or "")
            if "até" in texto or "ate" in texto:
                m = re.search(r"(\d{2})/(\d{2})/(\d{2,4})", texto)
                if m:
                    dia, mes, ano = m.groups()
                    ano = int(ano) if len(ano) == 4 else 2000 + int(ano)
                    return int(mes), ano
                return None
    return None


def mes_ano_do_nome(nome):
    m = re.match(r"Realizado_([A-Za-zç]+)_(\d{4})", nome)
    if m and m.group(1) in INDICE_MES:
        return INDICE_MES[m.group(1)], int(m.group(2))
    return None


def main():
    por_mes = {}
    for caminho in glob.glob(os.path.join(PASTA_REALIZADO, "Realizado_*.xlsx")):
        base = os.path.basename(caminho)
        if base.startswith("~$"):
            continue  # lock/temp do Excel
        ws = openpyxl.load_workbook(caminho, data_only=True).worksheets[0]
        chave = mes_ano_do_periodo(ws) or mes_ano_do_nome(base)
        if not chave:
            continue
        col_cred, col_deb = colunas_credito_debito(ws)
        credito = debito = 0.0
        for linha in ws.iter_rows(values_only=True):
            try:
                cred, deb = float(linha[col_cred] or 0), float(linha[col_deb] or 0)
            except (TypeError, ValueError, IndexError):
                continue  # cabecalho / totais / texto
            credito += cred
            debito += deb
        # Se dois arquivos mapearem para o mesmo mes, fica o baixado por ultimo
        existente = por_mes.get(chave)
        if existente is None or os.path.getmtime(caminho) >= existente[0]:
            por_mes[chave] = (os.path.getmtime(caminho),
                              {"mes": f"{ORDEM[chave[0] - 1]}/{chave[1] % 100:02d}",
                               "entradas": round(credito, 2),
                               "saidas": round(debito, 2)})

    ano_atual = date.today().year
    registros = [
        r for (_, r) in por_mes.values()
        if 2000 + int(r["mes"].split("/")[1]) == ano_atual
        and (r["entradas"] > 0 or r["saidas"] > 0)
    ]
    registros.sort(key=lambda r: (int(r["mes"].split("/")[1]),
                                  ORDEM.index(r["mes"].split("/")[0])))

    with open(SAIDA, "w", encoding="utf-8") as f:
        json.dump(registros, f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"{len(registros)} meses gravados em {SAIDA}")


if __name__ == "__main__":
    main()
