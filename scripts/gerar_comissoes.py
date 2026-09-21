"""
Gera src/data/comissoes.json a partir do relatório Focco de contratos
(loja_18314_MOEMA_CONTRATOS/relatorios_focco), usando as colunas
"Valor da Venda" e duas datas-base:

  1) "Assinatura"          -> séries "series"        (visão por assinatura)
  2) "Previsão de Entrega" -> séries "seriesEntrega" (visão por entrega)

Regras aplicadas sobre a data-base escolhida:

  - Montador:  10% do valor da venda, lançado 65 dias após a data-base
  - Liberador:  2% do valor da venda, lançado no mês subsequente à data-base
  - Vendedor:   6% do valor da venda, lançado no mês subsequente à data-base
  - Gerente:    5% do valor da venda, lançado no mês subsequente à data-base
  - Medidor:    R$ 3.960,00 fixo, lançado em todos os meses da linha do tempo

Alimenta a aba "Comissões" — complementar à aba Fluxo de Caixa (/fluxo-caixa).
"""
import json
import os
import sys
from datetime import datetime, timedelta

from gerar_projecao_contratos import (
    CONTRATOS_DIR,
    parse_data,
    parse_valor,
    rotulo_mes,
    xlsx_mais_recente,
)

import openpyxl

# Regras de cálculo (percentual sobre o Valor da Venda + momento do lançamento)
TAXA_MONTADOR = 0.10
DEFASAGEM_MONTADOR_DIAS = 65
TAXA_LIBERADOR = 0.02
TAXA_VENDEDOR = 0.06
TAXA_GERENTE = 0.05
VALOR_MEDIDOR_MENSAL = 3960.0

CAMPOS_SERIE = ('montador', 'liberador', 'vendedor', 'gerente', 'medidor', 'total')


def achar_cabecalho(ws):
    """Linha do cabeçalho da tabela de contratos (colunas Valor da Venda,
    Assinatura, Previsão de Entrega...). O relatório Focco traz filtros antes."""
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        celulas = [str(c).strip() if c is not None else '' for c in row]
        if ('Valor da Venda' in celulas and 'Assinatura' in celulas
                and 'Previsão de Entrega' in celulas):
            return i, {nome: j for j, nome in enumerate(celulas) if nome}
    raise ValueError('Cabeçalho "Valor da Venda"/"Assinatura"/'
                     '"Previsão de Entrega" não encontrado')


def mes_subsequente(d):
    """(ano, mês) do mês seguinte à data d."""
    return (d.year + 1, 1) if d.month == 12 else (d.year, d.month + 1)


def proxima_chave(chave):
    """Itera a faixa contínua de meses: (ano, mes) -> próximo (ano, mes)."""
    ano, mes = chave
    return (ano + 1, 1) if mes == 12 else (ano, mes + 1)


def calcular_series(ws, linha_header, idx_valor, idx_base, nome_data_base):
    """Série mensal de comissões tomando a coluna idx_base (Assinatura ou
    Previsão de Entrega) como data-base dos lançamentos. Devolve
    (series, n_contratos, total_base, n_sem_data)."""
    por_mes = {}  # (ano, mes) -> lançamentos do mês
    n_contratos = 0
    total_base = 0.0
    n_sem_data = 0

    def bucket(chave):
        return por_mes.setdefault(
            chave,
            {'montador': 0.0, 'liberador': 0.0, 'vendedor': 0.0, 'gerente': 0.0,
             'contratos': 0, 'valorBase': 0.0})

    for row in ws.iter_rows(min_row=linha_header + 1, values_only=True):
        base = parse_data(row[idx_base] if idx_base < len(row) else None)
        valor = parse_valor(row[idx_valor] if idx_valor < len(row) else None)
        if base is None:
            if valor:
                n_sem_data += 1
            continue

        reg = bucket((base.year, base.month))
        reg['contratos'] += 1
        reg['valorBase'] += valor

        # Liberador, Vendedor e Gerente: mês subsequente à data-base
        subsequente = bucket(mes_subsequente(base))
        subsequente['liberador'] += valor * TAXA_LIBERADOR
        subsequente['vendedor'] += valor * TAXA_VENDEDOR
        subsequente['gerente'] += valor * TAXA_GERENTE

        # Montador: 65 dias após a data-base
        montagem = base + timedelta(days=DEFASAGEM_MONTADOR_DIAS)
        bucket((montagem.year, montagem.month))['montador'] += valor * TAXA_MONTADOR

        n_contratos += 1
        total_base += valor

    if not por_mes:
        raise ValueError(f'Nenhum contrato com data em "{nome_data_base}"')

    # Faixa contínua de meses (do primeiro ao último lançamento); o Medidor é
    # um custo fixo presente em todos eles.
    series = []
    chave = min(por_mes)
    fim = max(por_mes)
    while True:
        agg = por_mes.get(chave, {'montador': 0.0, 'liberador': 0.0, 'vendedor': 0.0,
                                  'gerente': 0.0, 'contratos': 0, 'valorBase': 0.0})
        total = (agg['montador'] + agg['liberador'] + agg['vendedor']
                 + agg['gerente'] + VALOR_MEDIDOR_MENSAL)
        series.append({
            'mes': rotulo_mes(datetime(chave[0], chave[1], 1)),
            'contratos': agg['contratos'],
            'valorBase': round(agg['valorBase'], 2),
            'montador': round(agg['montador'], 2),
            'liberador': round(agg['liberador'], 2),
            'vendedor': round(agg['vendedor'], 2),
            'gerente': round(agg['gerente'], 2),
            'medidor': VALOR_MEDIDOR_MENSAL,
            'total': round(total, 2),
        })
        if chave == fim:
            break
        chave = proxima_chave(chave)

    return series, n_contratos, total_base, n_sem_data


def renomear_base(series, novo_nome):
    """valorBase -> valorAssinado / valorEntregue conforme a data-base."""
    out = []
    for m in series:
        m2 = dict(m)
        m2[novo_nome] = m2.pop('valorBase')
        out.append(m2)
    return out


def somar(series, campos=CAMPOS_SERIE):
    return {c: round(sum(m[c] for m in series), 2) for c in campos}


def main():
    caminho = xlsx_mais_recente(CONTRATOS_DIR)
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb.worksheets[0]

    linha_header, colunas = achar_cabecalho(ws)
    idx_valor = colunas['Valor da Venda']

    # 1) Data-base: Assinatura
    series_assin, n_assin, total_assin, sem_assin = calcular_series(
        ws, linha_header, idx_valor, colunas['Assinatura'], 'Assinatura')
    series_assin = renomear_base(series_assin, 'valorAssinado')

    # 2) Data-base: Previsão de Entrega
    series_entrega, n_entrega, total_entrega, sem_entrega = calcular_series(
        ws, linha_header, idx_valor, colunas['Previsão de Entrega'],
        'Previsão de Entrega')
    series_entrega = renomear_base(series_entrega, 'valorEntregue')

    saida = {
        'fonte': os.path.basename(caminho),
        'colunas': ['Valor da Venda', 'Assinatura', 'Previsão de Entrega'],
        'regras': {
            'montador': {'taxa': TAXA_MONTADOR, 'defasagemDias': DEFASAGEM_MONTADOR_DIAS},
            'liberador': {'taxa': TAXA_LIBERADOR, 'defasagemMeses': 1},
            'vendedor': {'taxa': TAXA_VENDEDOR, 'defasagemMeses': 1},
            'gerente': {'taxa': TAXA_GERENTE, 'defasagemMeses': 1},
            'medidor': {'valorMensal': VALOR_MEDIDOR_MENSAL},
        },
        # Visão por Assinatura
        'totalAssinado': round(total_assin, 2),
        'numContratos': n_assin,
        'semAssinatura': sem_assin,
        'mesInicio': series_assin[0]['mes'],
        'mesFim': series_assin[-1]['mes'],
        'totais': somar(series_assin),
        'series': series_assin,
        # Visão por Previsão de Entrega (mesmas regras, data-base trocada)
        'totalEntregue': round(total_entrega, 2),
        'numContratosEntrega': n_entrega,
        'semPrevisaoEntrega': sem_entrega,
        'mesInicioEntrega': series_entrega[0]['mes'],
        'mesFimEntrega': series_entrega[-1]['mes'],
        'totaisEntrega': somar(series_entrega),
        'seriesEntrega': series_entrega,
    }

    destino = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                           'src', 'data', 'comissoes.json')
    with open(destino, 'w', encoding='utf-8') as f:
        json.dump(saida, f, ensure_ascii=False, indent=2)
        f.write('\n')

    fmt = lambda v: f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

    def resumo(rotulo, series, n, total, totais, sem_data):
        print(f'{rotulo}: {len(series)} meses ({series[0]["mes"]}-{series[-1]["mes"]}), '
              f'{n} contratos, base R$ {fmt(total)}')
        print('  ' + ' | '.join(f'{f}: R$ {fmt(totais[f])}'
                                for f in ('montador', 'liberador', 'vendedor',
                                          'gerente', 'medidor')))
        print(f'  total projetado: R$ {fmt(totais["total"])}')
        if sem_data:
            print(f'  aviso: {sem_data} linhas com valor mas sem data-base')

    resumo('comissoes (Assinatura)', series_assin, n_assin, total_assin,
           saida['totais'], sem_assin)
    resumo('comissoes (Previsão de Entrega)', series_entrega, n_entrega,
           total_entrega, saida['totaisEntrega'], sem_entrega)

    return 0


if __name__ == '__main__':
    sys.exit(main())
